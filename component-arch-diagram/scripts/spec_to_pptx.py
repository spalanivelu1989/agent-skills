"""Turn a component-arch-diagram spec into an editable PowerPoint slide.

Every band, group, box, label and arrow becomes a native PowerPoint shape, placed
at the same coordinates as the SVG (scaled), so the slide matches the PNG and can
be edited box by box.

    python spec_to_pptx.py spec.json out.pptx
"""
import json
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from stack import load_spec  # noqa: E402

DEFAULT_KINDS = {
    "ui": ("#EEF4FF", "#6B8BC9"), "api": ("#F1F5F9", "#64748B"), "agent": ("#FFF7E8", "#C8913A"),
    "svc": ("#F3EEFF", "#8C6BC9"), "data": ("#EAF4EC", "#5B8F66"), "model": ("#EAF7F7", "#3F9A9A"),
    "eval": ("#FFF0F3", "#C9546E"), "ops": ("#FFF8E1", "#B8962E"), "ext": ("#2F2F33", "#111111"),
}
LINE = "#56606E"
FONT = "Arial"

spec = load_spec(sys.argv[1])   # applies band "tone"s
kinds = {k: v for k, v in DEFAULT_KINDS.items()}
for k, v in spec.get("kinds", {}).items():
    kinds[k] = (v["fill"], v["stroke"])
W, H = spec["canvas"]["width"], spec["canvas"]["height"]

prs = Presentation()
prs.slide_width = Emu(12192000)                 # 13.333 in
S = prs.slide_width / W                          # EMU per canvas px
prs.slide_height = Emu(int(H * S))
slide = prs.slides.add_slide(prs.slide_layouts[6])
PT = S / 12700                                   # points per canvas px (for fonts)


def rgb(h):
    return RGBColor.from_string(h.lstrip("#"))


def E(v):
    return Emu(int(round(v * S)))


def style(shape, fill, stroke, width=1.3):
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    if stroke is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(stroke)
        shape.line.width = Pt(width * PT)
    shape.shadow.inherit = False
    st = shape._element.find(qn("p:style"))   # theme style adds a shadow in some viewers
    if st is not None:
        shape._element.remove(st)


def radius(shape, px, w, h):
    try:
        shape.adjustments[0] = min(0.5, px / min(w, h))
    except IndexError:
        pass


def text(shape, runs, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=4):
    """runs: list of paragraphs, each (text, size_px, bold, color, letter_spacing_px)."""
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, E(margin))
    for i, (t, size, bold, color, spacing) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = t
        r.font.size = Pt(size * PT)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = rgb(color)
        if spacing:
            r._r.get_or_add_rPr().set("spc", str(int(spacing * PT * 100)))


def textbox(x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    text(tb, runs, align, anchor, margin=0)
    return tb


# ---- title
if spec.get("title"):
    textbox(0, 12, W, 40, [(spec["title"], 24, True, "#1E2530", 0)], PP_ALIGN.CENTER)

# ---- bands and side groups (behind everything)
for b in spec.get("bands", []):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(b["x"]), E(b["y"]), E(b["w"]), E(b["h"]))
    style(sh, b.get("fill", "#FBFCFE"), b.get("stroke", "#C9D1DB"))
    radius(sh, b.get("rx", 14), b["w"], b["h"])
    sh.name = f"Band: {b['name']}"
    if b.get("caps"):
        textbox(b["x"] + 28, b["y"] + 16, 700, 28,
                [(b["name"].upper(), 19, True, b.get("title_color", "#1E2530"), 3.5)])
        if b.get("desc"):
            textbox(b["x"] + 28, b["y"] + 46, b["w"] - 56, 22, [(b["desc"], 14, False, b.get("desc_color", "#4A5563"), 0)])
    elif b.get("tinted"):
        textbox(b["x"] + 52, b["y"] + 20, b["w"] - 104, 36, [(b["name"], 27, True, b.get("title_color", "#1E2530"), 0)])
        if b.get("desc"):
            textbox(b["x"] + 52, b["y"] + 62, b["w"] - 104, 30, [(b["desc"], 21, False, b.get("desc_color", "#4A5563"), 0)])
    else:
        textbox(b["x"] + 20, b["y"] + 16, 160, 60, [(b.get("n", ""), 13, True, "#5B6B7F", 0)] +
                [(part, 15, True, "#1E2530", 0) for part in b["name"].split("\n")])
for g in spec.get("groups", []):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(g["x"]), E(g["y"]), E(g["w"]), E(g["h"]))
    style(sh, "#FBFCFE", "#C9D1DB")
    radius(sh, 14, g["w"], g["h"])
    sh.name = f"Group: {g['name']}"
    label = (g["n"] + "  " if g.get("n") else "") + g["name"]
    textbox(g["x"] + 16, g["y"] + 10, g["w"] - 32, 24, [(label, 15, True, "#1E2530", 0)])


# ---- links (under the boxes): one connector per segment, arrowhead on the last
def arrowhead(line_shape):
    ln = line_shape.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "arrow" if STYLE.get("head") == "open" else "triangle")
    tail.set("w", "med")
    tail.set("len", "med")


STYLE = spec.get("link_style", {})   # see render.py: open heads, line colour
for i, lk in enumerate(spec["links"]):
    pts = lk["pts"]
    color = lk.get("color", STYLE.get("color", LINE))
    segs = list(zip(pts, pts[1:]))
    for j, (p, q) in enumerate(segs):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(p[0]), E(p[1]), E(q[0]), E(q[1]))
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(STYLE.get("width", 1.7) * PT)
        if lk.get("dashed"):
            c.line.dash_style = 4          # dash
        if j == len(segs) - 1 and lk.get("arrow", True):
            arrowhead(c)
        c.name = f"Link {i}" + (f": {lk['label']}" if lk.get("label") else "")
    if lk.get("label"):
        x, y = lk["label_at"]
        anchor = lk.get("anchor", "start")
        w = 7 * len(lk["label"]) + 10
        x0 = x - (w / 2 if anchor == "middle" else w if anchor == "end" else 0)
        textbox(x0, y - 13, w, 18, [(lk["label"], 12, False, color if lk.get("color") else "#37424F", 0)],
                {"start": PP_ALIGN.LEFT, "middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}[anchor])

# ---- actors
for a in spec.get("actors", []):
    x, y = a["cx"], a["cy"] - 25
    col = a.get("color", "#3B4A5E") if a.get("style") == "bust" else LINE
    head = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(x - 12), E(y - 8), E(24), E(24))
    style(head, col, None)
    body = slide.shapes.add_shape(MSO_SHAPE.ROUND_2_SAME_RECTANGLE, E(x - 22), E(y + 19), E(44), E(25))
    style(body, col, None)
    body.adjustments[0] = 0.5
    head.name, body.name = "User (head)", "User (body)"
    if a.get("label"):
        textbox(x + 35, y + 8, 160, 40, [(t, 14, True, "#1E2530", 0) for t in a["label"].split("\n")])

# ---- boxes
SHAPES = {"rect": MSO_SHAPE.ROUNDED_RECTANGLE, "cyl": MSO_SHAPE.CAN, "folder": MSO_SHAPE.FOLDED_CORNER,
          "cloud": MSO_SHAPE.ROUNDED_RECTANGLE, "hex": MSO_SHAPE.HEXAGON}
box = spec["box"]
bw, bh = box["w"], box["h"]
ts, ss, brx = box.get("title_size", 15), box.get("sub_size", 12.5), box.get("rx", 10)
for n in spec["nodes"]:
    w, h = n.get("w", bw), n.get("h", bh)
    shape = n.get("shape", "rect")
    fill, stroke = kinds[n["kind"]]
    dark = n["kind"] == "ext" or shape == "cloud"
    if shape == "cloud":
        fill, stroke = "#2F2F33", "#111111"
    sh = slide.shapes.add_shape(SHAPES[shape], E(n["cx"] - w / 2), E(n["cy"] - h / 2), E(w), E(h))
    style(sh, fill, stroke, 1.5)
    if shape in ("rect", "cloud"):
        radius(sh, 30 if shape == "cloud" else brx, w, h)
    if shape == "cyl":
        sh.adjustments[0] = 0.12
    sh.name = n["title"]
    tc, sc = ("#FFFFFF", "#D6D6DA") if dark else ("#1E2530", "#4A5563")
    runs = [(n["title"], ts, True, tc, 0)] + [(l, ss, False, sc, 0) for l in n.get("sub", "").split("\n") if l]
    text(sh, runs)

# ---- legend
lg = spec.get("legend")
if lg:
    b = lg["box"]
    lines = lg.get("lines", [])
    hgt = 22 + 18 * len(lines)
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(b["x"]), E(b["y"]), E(b["w"]), E(hgt))
    style(sh, "#F8FAFC", "#D5DCE4", 1)
    radius(sh, 10, b["w"], hgt)
    sh.name = "Legend"
    text(sh, [(t, 12.5, False, "#4A5563", 0) for t in lines], PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE, margin=16)

prs.save(sys.argv[2])
print("wrote", sys.argv[2], len(slide.shapes), "shapes")
