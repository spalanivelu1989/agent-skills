#!/usr/bin/env python3
"""The same tech-stack diagram as an editable PowerPoint slide.

    python3 stack_to_pptx.py spec.json out.pptx      (needs python-pptx)

Every layer, card, label and arrow is a native shape at the same position as in
the SVG, so the slide can be edited card by card.
"""
import json
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from stack_diagram import (CARD, DARK, INK, LAYER_DESC, LAYER_TITLE, LINE, MUTED,  # noqa: E402
                           SUB_SIZE, TITLE_SIZE, layout)

FONT = "Arial"
lay = layout(json.load(open(sys.argv[1])))
W, H = lay["W"], lay["H"]
prs = Presentation()
prs.slide_width = Emu(12192000)                 # 13.333 in wide; height follows the diagram
S = prs.slide_width / W
prs.slide_height = Emu(int(H * S))
slide = prs.slides.add_slide(prs.slide_layouts[6])
PT = S / 12700                                   # points per canvas px


def rgb(h):
    return RGBColor.from_string(h.lstrip("#"))


def E(v):
    return Emu(int(round(v * S)))


def rounded(x, y, w, h, radius, fill, stroke, width):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, E(x), E(y), E(w), E(h))
    sh.adjustments[0] = min(0.5, radius / min(w, h))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    sh.line.color.rgb = rgb(stroke); sh.line.width = Pt(width * PT)
    sh.shadow.inherit = False
    return sh


def text(x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    for i, (t, size, bold, color) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run(); r.text = t
        r.font.size = Pt(size * PT); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = rgb(color)
    return tb


if lay["title"]:
    text(0, 14, W, 34, [(lay["title"], 24, True, INK)], PP_ALIGN.CENTER)
for L in lay["layers"]:
    rounded(L["x"], L["y"], L["w"], L["h"], 26, L["fill"], L["stroke"], 1.3).name = f"Layer: {L['name']}"
    text(L["x"] + 52, L["y"] + 22, L["w"] - 104, 36, [(L["name"], LAYER_TITLE, True, L["title_color"])])
    if L["desc"]:
        text(L["x"] + 52, L["y"] + 64, L["w"] - 104, 30, [(L["desc"], LAYER_DESC, False, L["desc_color"])])
for a in lay["arrows"]:
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(a["x"]), E(a["y0"]), E(a["x"]), E(a["y1"]))
    c.line.color.rgb = rgb(LINE); c.line.width = Pt(1.7 * PT)
    tail = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
    tail.set("type", "triangle"); tail.set("w", "med"); tail.set("len", "med")
    if a["label"]:
        text(a["x"] + 12, (a["y0"] + a["y1"]) / 2 - 9, 300, 20, [(a["label"], 14, False, "#37424F")])
for cd in lay["cards"]:
    st = DARK if cd["dark"] else {**CARD, "title": INK, "sub": MUTED}
    sh = rounded(cd["x"], cd["y"], cd["w"], cd["h"], 14, st["fill"], st["stroke"], 1.5)
    sh.name = f"Card: {cd['title']}"
    tf = sh.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    runs = [(cd["title"], TITLE_SIZE, True, st["title"])] + [(s, SUB_SIZE, False, st["sub"]) for s in cd["sub"].split("\n") if s]
    for i, (t, size, bold, color) in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = t
        r.font.size = Pt(size * PT); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = rgb(color)
lg = lay["legend"]
if lg:
    rounded(lg["x"], lg["y"], lg["w"], lg["h"], 10, "#F8FAFC", "#D5DCE4", 1)
    text(lg["x"] + 20, lg["y"] + 8, lg["w"] - 40, lg["h"] - 12, [(t, 14, False, MUTED) for t in lg["lines"]])

prs.save(sys.argv[2])
print(f"wrote {sys.argv[2]} — {len(slide.shapes)} shapes")
